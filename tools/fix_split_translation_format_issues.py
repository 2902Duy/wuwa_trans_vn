import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("mistral_translate_work/split_by_prompt")
JSON_ROOT = ROOT / "json"
EXCEL_ROOT = ROOT / "excel"
REPORT_ROOT = Path("mistral_translate_work/reports/fix_split_translation_format_issues")


EXACT_UI = {
    "Remove": "Gỡ bỏ",
    "Place": "Đặt",
    "Previous": "Trước",
    "Ongoing": "Đang diễn ra",
    "Requirement": "Yêu cầu",
    "Equipped": "Đã trang bị",
    "Unlocked": "Đã mở khóa",
    "Claim": "Nhận",
    "Revive": "Hồi sinh",
    "Controls": "Điều khiển",
    "Skills": "Kỹ năng",
    "Wave": "Đợt",
    "Leave": "Rời khỏi",
    "Rating": "Đánh giá",
    "Tactics": "Chiến thuật",
    "Confirm": "Xác nhận",
    "Details": "Chi tiết",
    "Locked": "Đã khóa",
    "Unavailable": "Không khả dụng",
    "Defeated": "Đã thất bại",
    "Track": "Theo dõi",
    "Team": "Đội",
    "Close": "Đóng",
    "Cancel": "Hủy",
    "Rewards": "Phần thưởng",
    "Challenge": "Thử thách",
    "Start": "Bắt đầu",
    "Continue": "Tiếp tục",
    "Retry": "Thử lại",
    "Skip": "Bỏ qua",
    "Next": "Tiếp",
    "Back": "Quay lại",
    "View": "Xem",
    "Obtained": "Đã nhận",
    "Unlock": "Mở khóa",
}


PHRASE_REPLACEMENTS = [
    (re.compile(r"\bHello\.\.\."), "Xin chào..."),
    (re.compile(r"\bTruly\?!"), "Thật sao?!"),
    (re.compile(r"\bTruly\?"), "Thật sao?"),
    (re.compile(r"\bWho\?"), "Ai?"),
    (re.compile(r"\bWhy\?"), "Tại sao?"),
    (re.compile(r"\bWhat\?!"), "Gì cơ?!"),
    (re.compile(r"\bWhat\?"), "Gì cơ?"),
    (re.compile(r"\bPlushies\?"), "Thú bông?"),
    (re.compile(r"\bUtility: Flight\b"), "Công Cụ: Bay Lượn"),
    (re.compile(r"\bUnion Level\b"), "Cấp Liên Minh"),
    (re.compile(r"\bUnion Levels\b"), "Cấp Liên Minh"),
    (re.compile(r"\bEvery 10 Cấp Liên Minh raises SOL3 Phase by 1\."), "Mỗi 10 Cấp Liên Minh tăng Pha SOL3 lên 1."),
    (re.compile(r"\b20, 40, and 60\b"), "20, 40 và 60"),
    (re.compile(r"\bPhase Ascension\b"), "Thăng Tiến Pha"),
    (re.compile(r"\bSupportive Pilgrim\b"), "Lữ Khách Hỗ Trợ"),
    (re.compile(r"\bIndignant Supporter\b"), "Người Ủng Hộ Phẫn Nộ"),
    (re.compile(r"\bNo Supported Cube\b"), "Không Có Cube Được Hỗ Trợ"),
    (re.compile(r"\bCube Supported\b"), "Cube Được Hỗ Trợ"),
    (re.compile(r"\bSupporting Cube\b"), "Cube Hỗ Trợ"),
    (re.compile(r"\bSupported\b"), "Được Hỗ Trợ"),
    (re.compile(r"\bSupporting\b"), "Đang Hỗ Trợ"),
    (re.compile(r"\bReset button\b"), "nút Đặt lại"),
    (re.compile(r"\bReset\b"), "Đặt lại"),
    (re.compile(r"\bSupport Echoes\b"), "Echo Hỗ Trợ"),
    (re.compile(r"\bSupport Echo\b"), "Echo Hỗ Trợ"),
    (re.compile(r"\bEcho Support\b"), "Echo Hỗ Trợ"),
    (re.compile(r"\bSupport Resonators\b"), "Resonator Hỗ Trợ"),
    (re.compile(r"\bSupport Resonator\b"), "Resonator Hỗ Trợ"),
    (re.compile(r"\bSupport Module\b"), "Mô-đun Hỗ Trợ"),
    (re.compile(r"\bSupport Item\b"), "Vật Phẩm Hỗ Trợ"),
    (re.compile(r"\bSupport Items\b"), "Vật Phẩm Hỗ Trợ"),
    (re.compile(r"\bSupport\b"), "Hỗ Trợ"),
    (re.compile(r"\bThank you\b"), "Cảm ơn"),
    (re.compile(r"\bThanks\b"), "Cảm ơn"),
    (re.compile(r"\bBut\b"), "Nhưng"),
    (re.compile(r"\bShe\b"), "Cô ấy"),
    (re.compile(r"\bHe\b"), "Anh ấy"),
    (re.compile(r"\bThey\b"), "Họ"),
    (re.compile(r"\bReal\b"), "Thật sự"),
    (re.compile(r"\bYou\b"), "Bạn"),
    (re.compile(r"\bSelect\b"), "Chọn"),
    (re.compile(r"\bTap\b"), "Chạm"),
    (re.compile(r"\bClaim\b"), "Nhận"),
    (re.compile(r"\bConfirm\b"), "Xác nhận"),
    (re.compile(r"\bResume\b"), "Tiếp tục"),
    (re.compile(r"\bRevive\b"), "Hồi sinh"),
    (re.compile(r"\bInsufficient materials\b"), "Không đủ nguyên liệu"),
    (re.compile(r"\bProceed to unlock\b"), "Tiến hành mở khóa"),
    (re.compile(r"\bCollect\b"), "Thu thập"),
    (re.compile(r"\bUpgrade\b"), "Nâng cấp"),
    (re.compile(r"\bHold\b"), "Giữ"),
    (re.compile(r"\bClick\b"), "Nhấp"),
    (re.compile(r"\bDrag\b"), "Kéo"),
    (re.compile(r"\bUse\b"), "Dùng"),
    (re.compile(r"\bOpen\b"), "Mở"),
    (re.compile(r"\bContinue\b"), "Tiếp tục"),
    (re.compile(r"\bWait\b"), "Khoan"),
    (re.compile(r"\bBack\b"), "Quay lại"),
    (re.compile(r"\bAdjust\b"), "Điều chỉnh"),
    (re.compile(r"\bThen\b"), "Vậy thì"),
    (re.compile(r"\bAfter\b"), "Sau khi"),
    (re.compile(r"\bToday\b"), "Hôm nay"),
    (re.compile(r"\bDay\b"), "Ngày"),
    (re.compile(r"\bReporter\b"), "Phóng viên"),
    (re.compile(r"\bItems\b"), "Vật phẩm"),
    (re.compile(r"\bDescription\b"), "Mô tả"),
    (re.compile(r"\bVoice\b"), "Giọng nói"),
    (re.compile(r"\bAlert\b"), "Cảnh báo"),
    (re.compile(r"\bModule\b"), "Mô-đun"),
    (re.compile(r"\bRequirement\b"), "Yêu cầu"),
    (re.compile(r"\bStandard\b"), "Tiêu Chuẩn"),
    (re.compile(r"\bExtra\b"), "Bổ Sung"),
    (re.compile(r"\bGameplay\b"), "Lối chơi"),
    (re.compile(r"\bMain Quest\b"), "Nhiệm Vụ Chính"),
    (re.compile(r"\bChapter\b"), "Chương"),
    (re.compile(r"\bExcuse me\b"), "Xin lỗi"),
    (re.compile(r"\bSo what\b"), "Vậy thì sao"),
    (re.compile(r"\bSalt\b"), "Muối"),
    (re.compile(r"\bOthers\b"), "Khác"),
    (re.compile(r"\bSword\b"), "Kiếm"),
    (re.compile(r"\bOperate\b"), "Vận hành"),
    (re.compile(r"\bResearch\b"), "Nghiên cứu"),
    (re.compile(r"\bNote\b"), "Ghi chú"),
    (re.compile(r"\bTargeted Merge\b"), "Hợp Nhất Chỉ Định"),
    (re.compile(r"\bTarget\b"), "Mục tiêu"),
    (re.compile(r"\bAuto Select\b"), "Tự Động Chọn"),
]


MOJIBAKE_MARKERS = ("Ã", "Â", "áº", "á»", "Ä", "Æ", "â€", "â€™", "â€œ", "â€�")


def maybe_fix_mojibake(text: str) -> str:
    if not any(marker in text for marker in MOJIBAKE_MARKERS):
        return text
    candidates = []
    for encoding in ("latin1", "cp1252"):
        try:
            candidates.append(text.encode(encoding).decode("utf-8"))
        except UnicodeError:
            pass
    if not candidates:
        return text

    def score(value: str) -> int:
        bad = sum(value.count(marker) for marker in MOJIBAKE_MARKERS)
        good = sum(value.count(ch) for ch in "ăâđêôơưĂÂĐÊÔƠƯàáạảãầấậẩẫằắặẳẵèéẹẻẽềếệểễìíịỉĩòóọỏõồốộổỗờớợởỡùúụủũừứựửữỳýỵỷỹ")
        return good * 4 - bad * 8 - value.count("?") * 2

    best = max(candidates, key=score)
    return best if score(best) > score(text) else text


def fix_text(text: str, domain: str) -> tuple[str, list[str]]:
    original = text
    reasons = []
    text = maybe_fix_mojibake(text)
    if text != original:
        reasons.append("mojibake")

    if domain == "ui" and text.strip() in EXACT_UI:
        new_text = EXACT_UI[text.strip()]
        if new_text != text:
            text = new_text
            reasons.append("exact_ui")
        return text, reasons

    before = text
    for pattern, replacement in PHRASE_REPLACEMENTS:
        text = pattern.sub(replacement, text)
    if text != before:
        reasons.append("mixed_english_verb")

    # Cleanup common artifacts after direct phrase replacement.
    cleanup = {
        "Bạn đã bị thêm": "Bạn đã được thêm",
        "Bạn có thể dùng": "Bạn có thể sử dụng",
        "Bạn có thể xem": "Bạn có thể xem",
        "Chạm để Start Game": "Chạm để bắt đầu trò chơi",
        "Chạm bất kỳ đâu": "Chạm vào bất kỳ đâu",
    }
    before = text
    for src, dst in cleanup.items():
        text = text.replace(src, dst)
    if text != before:
        reasons.append("cleanup")

    return text, reasons


def sync_excel(json_path: Path, updated_by_split_id: dict[str, str]) -> bool:
    rel = json_path.relative_to(JSON_ROOT)
    xlsx_path = EXCEL_ROOT / rel.with_suffix(".xlsx")
    if not xlsx_path.exists() or not updated_by_split_id:
        return False
    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    try:
        split_col = headers.index("split_id") + 1
        trans_col = headers.index("new_translation_vi") + 1
    except ValueError:
        return False
    changed = False
    for row in range(2, sheet.max_row + 1):
        split_id = sheet.cell(row=row, column=split_col).value
        if split_id in updated_by_split_id:
            sheet.cell(row=row, column=trans_col).value = updated_by_split_id[split_id]
            changed = True
    if changed:
        workbook.save(xlsx_path)
    return changed


def main() -> None:
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    changes = []
    file_counter = Counter()
    reason_counter = Counter()
    excel_synced = 0

    for json_path in JSON_ROOT.rglob("*.json"):
        rel = json_path.relative_to(JSON_ROOT)
        domain = rel.parts[0] if rel.parts else ""
        data = json.loads(json_path.read_text(encoding="utf-8"))
        if not isinstance(data, list):
            continue

        changed = False
        updated_by_split_id = {}
        for index, item in enumerate(data):
            if not isinstance(item, dict):
                continue
            value = item.get("new_translation_vi")
            if not isinstance(value, str) or not value:
                continue
            fixed, reasons = fix_text(value, domain)
            if fixed == value:
                continue
            item["new_translation_vi"] = fixed
            changed = True
            split_id = item.get("split_id") or f"{rel.as_posix()}#{index}"
            updated_by_split_id[split_id] = fixed
            file_counter[rel.as_posix()] += 1
            reason_counter.update(reasons)
            changes.append(
                {
                    "file": rel.as_posix(),
                    "index": index,
                    "split_id": split_id,
                    "domain": domain,
                    "reasons": reasons,
                    "before": value,
                    "after": fixed,
                    "source_en": item.get("source_en", ""),
                }
            )

        if changed:
            json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            if sync_excel(json_path, updated_by_split_id):
                excel_synced += 1

    summary = {
        "json_files_scanned": len(list(JSON_ROOT.rglob("*.json"))),
        "changes": len(changes),
        "excel_files_synced": excel_synced,
        "by_reason": dict(reason_counter.most_common()),
        "top_files": dict(file_counter.most_common(50)),
    }
    (REPORT_ROOT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (REPORT_ROOT / "changes.json").write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
