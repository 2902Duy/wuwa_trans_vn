import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("mistral_translate_work/split_by_prompt")
JSON_ROOT = ROOT / "json"
EXCEL_ROOT = ROOT / "excel"


def sync_one(json_path: Path) -> tuple[bool, int]:
    rel = json_path.relative_to(JSON_ROOT)
    xlsx_path = EXCEL_ROOT / rel.with_suffix(".xlsx")
    if not xlsx_path.exists():
        return False, 0

    data = json.loads(json_path.read_text(encoding="utf-8"), strict=False)
    if not isinstance(data, list):
        return False, 0
    by_split_id = {
        item.get("split_id"): item.get("new_translation_vi", "")
        for item in data
        if isinstance(item, dict) and item.get("split_id")
    }
    if not by_split_id:
        return False, 0

    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    try:
        split_col = headers.index("split_id") + 1
        trans_col = headers.index("new_translation_vi") + 1
    except ValueError:
        return False, 0

    changes = 0
    for row in range(2, sheet.max_row + 1):
        split_id = sheet.cell(row=row, column=split_col).value
        if split_id not in by_split_id:
            continue
        expected = by_split_id[split_id]
        current = sheet.cell(row=row, column=trans_col).value
        if current is None:
            current = ""
        if current != expected:
            sheet.cell(row=row, column=trans_col).value = expected
            changes += 1

    if changes:
        workbook.save(xlsx_path)
    return bool(changes), changes


def main() -> None:
    files = 0
    rows = 0
    for json_path in JSON_ROOT.rglob("*.json"):
        changed, count = sync_one(json_path)
        if changed:
            files += 1
            rows += count
    print({"excel_files_updated": files, "rows_updated": rows})


if __name__ == "__main__":
    main()
