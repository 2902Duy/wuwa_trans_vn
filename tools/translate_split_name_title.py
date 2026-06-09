import argparse
import json
import os
import re
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


WORKSPACE_DIR = Path(r"C:\Users\tduy2\Documents\antigravity\silly-darwin")
WORK_DIR = WORKSPACE_DIR / "mistral_translate_work"
PROMPT_DIR = WORK_DIR / "prompts_compressed"
NAME_TITLE_DIR = WORK_DIR / "split_by_prompt" / "json" / "name_title"
EXCEL_NAME_TITLE_DIR = WORK_DIR / "split_by_prompt" / "excel" / "name_title"
CLASSIFICATION_REPORT = WORK_DIR / "reports" / "split_name_title_classification" / "all_rows.json"
RUN_DIR = WORK_DIR / "reports" / "name_title_translation_run"

API_URL = "https://api.mistral.ai/v1/chat/completions"
MODEL_NAME = "mistral-large-latest"

TRANSLATE_CLASSES = {
    "translate_short_title",
    "translate_short_ui_title",
    "translate_role_or_speaker_title",
    "translate_ambiguous_name_title",
}

KEEP_CLASSES = {
    "keep_proper_name",
    "keep_item_weapon_echo_name",
    "keep_glossary_term",
    "keep_symbol_or_number",
    "keep_with_placeholder",
    "keep_placeholder_only",
}

EXCEL_COLUMNS = [
    "split_id",
    "prompt_domain",
    "prompt_file",
    "source_file",
    "original_index",
    "database",
    "table",
    "primary_key_column",
    "primary_key",
    "column",
    "category",
    "source_en",
    "new_translation_vi",
    "translator_note",
]

TOKEN_RE = re.compile(r"(<[^>]+>|\{\d+\}|\{[A-Za-z_][A-Za-z0-9_]*\}|\\[nrt]|\{Cus:[^}]+\})")
ROW_RE = re.compile(r"^\s*([A-Z_]+_\d{7}):::(.*)$")
GLOSSARY_TERMS = (
    "Rover",
    "Echo",
    "Resonator",
    "Resonance Chain",
    "Resonance Skill",
    "Resonance Liberation",
    "Forte Circuit",
    "Basic Attack",
    "Normal Attack",
    "Heavy Attack",
    "Mid-air Attack",
    "Dodge Counter",
    "Intro Skill",
    "Outro Skill",
    "Inherent Skill",
    "Aero Erosion",
    "Glacio Chafe",
    "Spectro Frazzle",
    "Havoc Bane",
    "Fusion Burst",
    "Electro Flare",
    "Negative Status",
    "Tidal Blight",
    "DMG",
    "DMG Bonus",
    "Crit. Rate",
    "Crit. DMG",
)



def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    content = json.dumps(data, ensure_ascii=False, indent=2) + "\n"
    for attempt in range(8):
        try:
            tmp.write_text(content, encoding="utf-8")
            # Using os.replace is atomic
            import os
            os.replace(str(tmp), str(path))
            return
        except PermissionError:
            if attempt == 7:
                path.write_text(content, encoding="utf-8")
                try:
                    tmp.unlink(missing_ok=True)
                except Exception:
                    pass
                return
            import time
            time.sleep(0.25 * (attempt + 1))


def normalize(text):
    return (text or "").replace("\r\n", "\\n").replace("\n", "\\n").replace("\r", "\\n")


def denormalize(text):
    return (text or "").replace("\\n", "\n")


def load_api_keys():
    keys = []
    env_path = WORKSPACE_DIR / ".env"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            name, value = line.split("=", 1)
            if name.strip() in {"MISTRAL_API_KEYS", "MISTRAL_API_KEY"}:
                keys.extend(part.strip() for part in value.strip().strip('"').strip("'").split(","))
    env_keys = os.environ.get("MISTRAL_API_KEYS") or os.environ.get("MISTRAL_API_KEY", "")
    if env_keys:
        keys.extend(part.strip() for part in env_keys.replace("\n", ",").split(","))

    clean = []
    for key in keys:
        if key and key not in clean:
            clean.append(key)
    return clean


def build_system_prompt():
    parts = []
    for name in ("shared_glossary.md", "keep_english_rules.md", "name_title_prompt.md"):
        parts.append((PROMPT_DIR / name).read_text(encoding="utf-8"))
    parts.append(
        """
## Strict Batch Output

Translate only rows that are meant to be translated. Preserve all IDs exactly.
Return only lines in this exact format:
ID:::Vietnamese translation

No Markdown, no comments, no explanations.
Keep all placeholders, tags, numbers, punctuation markers, and glossary terms exactly.
For ambiguous short titles, translate naturally as a game title while preserving proper nouns.
"""
    )
    return "\n\n".join(parts)


def build_classification_map():
    data = read_json(CLASSIFICATION_REPORT)["rows"]
    return {row["split_id"]: row["classification"] for row in data}


def iter_json_files():
    return sorted(NAME_TITLE_DIR.rglob("*.json"))


def set_keep_rows(classification):
    changed_files = 0
    changed_rows = 0
    for path in iter_json_files():
        rows = read_json(path)
        changed = False
        for row in rows:
            split_id = row["split_id"]
            cls = classification.get(split_id, "")
            if cls in KEEP_CLASSES:
                source = row.get("source_en", "")
                if row.get("new_translation_vi", "") != source:
                    row["new_translation_vi"] = source
                    row["translator_note"] = cls
                    changed = True
                    changed_rows += 1
        if changed:
            write_json(path, rows)
            changed_files += 1
    return changed_files, changed_rows


def collect_translate_items(classification, limit=None):
    items = []
    for path in iter_json_files():
        rows = read_json(path)
        for row in rows:
            cls = classification.get(row["split_id"], "")
            if cls not in TRANSLATE_CLASSES:
                continue
            if row.get("new_translation_vi"):
                continue
            item = {
                "json_path": str(path),
                "split_id": row["split_id"],
                "source_en": row.get("source_en", ""),
                "classification": cls,
            }
            items.append(item)
            if limit and len(items) >= limit:
                return items
    return items


def make_batches(items, batch_size, max_chars):
    batch = []
    chars = 0
    for item in items:
        row_len = len(item["source_en"]) + 40
        if batch and (len(batch) >= batch_size or chars + row_len > max_chars):
            yield batch
            batch = []
            chars = 0
        batch.append(item)
        chars += row_len
    if batch:
        yield batch


def build_user_prompt(batch):
    lines = []
    for item in batch:
        lines.append(f"{item['split_id']}:::{normalize(item['source_en'])}")
    return "\n".join(lines)


def call_mistral(api_key, system_prompt, user_prompt, retries):
    body = {
        "model": MODEL_NAME,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.1,
        "max_tokens": 8192,
    }
    data = json.dumps(body).encode("utf-8")
    delay = 5
    for attempt in range(1, retries + 1):
        request = urllib.request.Request(
            API_URL,
            data=data,
            method="POST",
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
                "Accept": "application/json",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=180) as response:
                payload = json.loads(response.read().decode("utf-8"))
                return payload["choices"][0]["message"]["content"]
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            if attempt < retries and exc.code in {429, 500, 502, 503, 504}:
                time.sleep(delay)
                delay *= 2
                continue
            raise RuntimeError(f"HTTP {exc.code}: {detail}") from exc
        except Exception:
            if attempt >= retries:
                raise
            time.sleep(delay)
            delay *= 2
    raise RuntimeError("Mistral request failed")


def call_mistral_key_pool(keys, start_index, system_prompt, user_prompt, retries, batch_delay):
    errors = []
    key_count = len(keys)
    current = start_index
    for round_no in range(1, retries + 1):
        for _ in range(key_count):
            key_index = current % key_count
            current += 1
            try:
                output = call_mistral(keys[key_index], system_prompt, user_prompt, 1)
                if batch_delay:
                    time.sleep(batch_delay)
                return output, current
            except Exception as exc:
                errors.append({"round": round_no, "key_index": key_index, "error": repr(exc)})
                continue
        if round_no < retries:
            time.sleep(15 * round_no)
    raise RuntimeError(json.dumps({"kind": "all_keys_failed", "errors": errors[-20:]}, ensure_ascii=False))


def parse_output(output):
    parsed = {}
    current_id = None
    for line in output.splitlines():
        match = ROW_RE.match(line)
        if match:
            current_id = match.group(1)
            parsed[current_id] = match.group(2).strip()
        elif current_id and line.strip():
            parsed[current_id] += "\n" + line.rstrip()
    return parsed


def validate_translation(source, translation):
    issues = []
    source_norm = normalize(source)
    translation_norm = normalize(translation)
    for token in TOKEN_RE.findall(source_norm):
        if token not in translation_norm:
            issues.append(f"missing_token:{token}")
    for term in GLOSSARY_TERMS:
        if term in source_norm and term not in translation_norm:
            issues.append(f"missing_glossary_term:{term}")
    return issues


def apply_translations(translations):
    by_path = defaultdict(dict)
    for split_id, translation in translations.items():
        by_path[translation["json_path"]][split_id] = translation

    changed_files = 0
    changed_rows = 0
    for path_str, by_id in by_path.items():
        path = Path(path_str)
        rows = read_json(path)
        changed = False
        for row in rows:
            item = by_id.get(row["split_id"])
            if not item:
                continue
            row["new_translation_vi"] = item["translation_vi"]
            row["translator_note"] = item["classification"]
            changed = True
            changed_rows += 1
        if changed:
            write_json(path, rows)
            changed_files += 1
    return changed_files, changed_rows


def write_excel(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "translation"
    ws.append(EXCEL_COLUMNS)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append([row.get(column, "") for column in EXCEL_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 20,
        "B": 18,
        "C": 28,
        "D": 28,
        "E": 14,
        "F": 24,
        "G": 24,
        "H": 18,
        "I": 38,
        "J": 16,
        "K": 14,
        "L": 80,
        "M": 80,
        "N": 36,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def regenerate_name_title_excels():
    count = 0
    for json_path in iter_json_files():
        rel = json_path.relative_to(NAME_TITLE_DIR)
        xlsx_path = EXCEL_NAME_TITLE_DIR / rel.with_suffix(".xlsx")
        write_excel(xlsx_path, read_json(json_path))
        count += 1
    return count


def count_completion(classification):
    counts = Counter()
    for path in iter_json_files():
        for row in read_json(path):
            cls = classification.get(row["split_id"], "")
            if cls in TRANSLATE_CLASSES:
                counts["translate_total"] += 1
                if row.get("new_translation_vi"):
                    counts["translate_done"] += 1
            elif cls in KEEP_CLASSES:
                counts["keep_total"] += 1
                if row.get("new_translation_vi") == row.get("source_en", ""):
                    counts["keep_done"] += 1
    return counts


def main():
    import concurrent.futures
    import threading

    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--batch-size", type=int, default=50)
    parser.add_argument("--max-chars", type=int, default=5000)
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--batch-delay", type=float, default=0.5)
    parser.add_argument("--workers", type=int, default=8)
    parser.add_argument("--skip-api", action="store_true")
    args = parser.parse_args()

    RUN_DIR.mkdir(parents=True, exist_ok=True)
    classification = build_classification_map()
    keep_files, keep_rows = set_keep_rows(classification)

    items = collect_translate_items(classification, limit=args.limit)
    keys = load_api_keys()
    if not keys and items and not args.skip_api:
        raise RuntimeError("No Mistral API keys found in .env or environment.")

    system_prompt = build_system_prompt()
    placeholder_errors = []
    translation_errors = []
    success = {}

    write_lock = threading.Lock()
    key_cursor_lock = threading.Lock()

    key_cursor = 0
    batch_counter = 0

    batches = list(make_batches(items, args.batch_size, args.max_chars))
    print(f"Total items to translate: {len(items)}, Batches: {len(batches)}, Parallel Workers: {args.workers}")

    def process_batch(batch):
        nonlocal key_cursor, batch_counter
        
        with write_lock:
            batch_counter += 1
            current_batch_no = batch_counter

        batch_ids = {item["split_id"] for item in batch}
        
        with key_cursor_lock:
            start_idx = key_cursor
            key_cursor += 1

        try:
            output, new_cursor = call_mistral_key_pool(
                keys,
                start_idx,
                system_prompt,
                build_user_prompt(batch),
                args.retries,
                args.batch_delay,
            )
            
            with key_cursor_lock:
                key_cursor = max(key_cursor, new_cursor)

            parsed = parse_output(output)
            missing_ids = sorted(batch_ids - set(parsed))
            
            local_placeholder_errors = []
            local_translation_errors = []
            batch_success = {}

            if missing_ids:
                local_translation_errors.append({
                    "kind": "missing_output_rows", 
                    "split_ids": missing_ids, 
                    "batch": current_batch_no
                })

            for item in batch:
                translated = parsed.get(item["split_id"])
                if not translated:
                    continue
                translated = denormalize(translated)
                issues = validate_translation(item["source_en"], translated)
                error_item = {
                    "split_id": item["split_id"],
                    "json_path": item["json_path"],
                    "classification": item["classification"],
                    "source_en": item["source_en"],
                    "translation_vi": translated,
                    "issues": issues,
                    "batch": current_batch_no,
                }
                if any(issue.startswith("missing_token:") for issue in issues):
                    local_placeholder_errors.append(error_item)
                    continue
                if issues:
                    local_translation_errors.append(error_item)
                    continue
                batch_success[item["split_id"]] = {
                    **item,
                    "translation_vi": translated,
                }

            with write_lock:
                if batch_success:
                    apply_translations(batch_success)
                    success.update(batch_success)
                    write_json(RUN_DIR / "latest_success_checkpoint.json", success)
                
                if local_placeholder_errors:
                    placeholder_errors.extend(local_placeholder_errors)
                if local_translation_errors:
                    translation_errors.extend(local_translation_errors)

                print(
                    json.dumps(
                        {
                            "batch": current_batch_no,
                            "rows": len(batch),
                            "success": len(batch_success),
                            "placeholder_errors": len(placeholder_errors),
                            "translation_errors": len(translation_errors),
                        },
                        ensure_ascii=False,
                    )
                )
        except Exception as exc:
            with write_lock:
                translation_errors.append(
                    {
                        "kind": "api_or_batch_error",
                        "batch": current_batch_no,
                        "split_ids": sorted(batch_ids),
                        "error": repr(exc),
                    }
                )
                write_json(RUN_DIR / "translation_errors.json", {"count": len(translation_errors), "errors": translation_errors})
            raise

    if not args.skip_api and batches:
        with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = [executor.submit(process_batch, b) for b in batches]
            for future in concurrent.futures.as_completed(futures):
                try:
                    future.result()
                except Exception as e:
                    print(f"Batch failed with exception: {e}")

    excel_count = regenerate_name_title_excels()
    completion = count_completion(classification)
    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "keep_rows_written": keep_rows,
        "keep_files_changed": keep_files,
        "api_translate_items_selected": len(items),
        "api_batches": len(batches),
        "api_success_rows": len(success),
        "placeholder_errors": len(placeholder_errors),
        "translation_errors": len(translation_errors),
        "excel_files_regenerated": excel_count,
        "completion": dict(completion),
    }
    write_json(RUN_DIR / "placeholder_tag_errors.json", {"count": len(placeholder_errors), "errors": placeholder_errors})
    write_json(RUN_DIR / "translation_errors.json", {"count": len(translation_errors), "errors": translation_errors})
    write_json(RUN_DIR / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
